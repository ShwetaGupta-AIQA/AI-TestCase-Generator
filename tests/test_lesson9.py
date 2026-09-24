import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from services.document_reader import read_document
from services.requirement_extractor import extract_requirements
from app import run_document_pipeline
from test_pipeline import ANALYSIS, SCENARIO, CASE


class Lesson9Tests(unittest.TestCase):
    def test_txt_docx_pdf(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            text = root / "BRD.TXT"
            text.write_text("REQ-1: Login with password", encoding="utf-8-sig")
            self.assertEqual(read_document(text), "REQ-1: Login with password")
            document = Document()
            document.add_paragraph("REQ-1: Login")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "REQ-2: Logout"
            document.add_paragraph("REQ-3: Dashboard")
            docx = root / "brd.docx"
            document.save(docx)
            self.assertEqual(read_document(docx).splitlines(), ["REQ-1: Login", "REQ-2: Logout", "REQ-3: Dashboard"])
            writer = PdfWriter()
            page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(b"BT /F1 12 Tf 20 250 Td (REQ-1: Login) Tj ET")
            page[NameObject("/Contents")] = stream
            pdf = root / "brd.pdf"
            writer.write(pdf)
            self.assertIn("REQ-1: Login", read_document(pdf))

    def test_reader_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, content in [("empty.txt", " "), ("bad.docx", "invalid"), ("bad.pdf", "invalid"), ("bad.csv", "data")]:
                path = root / name
                path.write_text(content)
                with self.subTest(name=name), self.assertRaises(ValueError):
                    read_document(path)
            with self.assertRaisesRegex(ValueError, "not found"):
                read_document(root / "missing.txt")
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=300)
            blank = root / "blank.pdf"
            writer.write(blank)
            with self.assertRaisesRegex(ValueError, "OCR"):
                read_document(blank)

    def test_extraction_and_first_only_pipeline(self):
        extracted = {"requirements": [{"source_id": "REQ-1", "requirement_text": "Login with password"}, {"source_id": "REQ-2", "requirement_text": "Redirect to dashboard"}]}
        with contextlib.ExitStack() as stack:
            stack.enter_context(contextlib.redirect_stderr(io.StringIO()))
            stack.enter_context(patch("services.requirement_extractor.call_llm", return_value="```json\n" + json.dumps(extracted) + "\n```"))
            analysis = stack.enter_context(patch("services.requirement_analyzer.call_llm", return_value=json.dumps(ANALYSIS)))
            stack.enter_context(patch("services.scenario_generator.call_llm", return_value=json.dumps({"scenarios": [SCENARIO]})))
            stack.enter_context(patch("services.test_case_generator.call_llm", return_value=json.dumps({"test_cases": [CASE]})))
            result = run_document_pipeline("sample_documents/login_brd.txt")
        analysis.assert_called_once()
        self.assertEqual(result["source_id"], "REQ-1")
        self.assertEqual(result["requirement"], "Login with password")
        self.assertEqual(result["analysis"]["requirement_id"], "REQ001")
        self.assertEqual(result["extracted_requirements"], extracted["requirements"])

    def test_no_requirements_stops_pipeline(self):
        with patch("services.requirement_extractor.call_llm", return_value='{"requirements": []}'), patch("app.run_pipeline") as pipeline:
            with self.assertRaisesRegex(ValueError, "No testable requirements"):
                run_document_pipeline("sample_documents/login_brd.txt")
            pipeline.assert_not_called()

    def test_empty_input_never_calls_model(self):
        with patch("services.requirement_extractor.call_llm") as call:
            with self.assertRaises(ValueError):
                extract_requirements(" ")
            call.assert_not_called()
