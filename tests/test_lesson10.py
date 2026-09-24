import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from document_pipeline import run_document_pipeline
from services.excel_exporter import export_to_excel
from models.document_requirement import ExtractedRequirement
from models.requirement import RequirementAnalysis
from models.test_scenario import TestScenario
from models.test_case import TestCase
from test_pipeline import ANALYSIS, SCENARIO, CASE


class Lesson10Tests(unittest.TestCase):
    def records(self):
        return ([{"source": ExtractedRequirement(source_id="REQ-2", requirement_text="Password length 8 to 20"),
                  "analysis": RequirementAnalysis(**ANALYSIS)}],
                [TestScenario(**SCENARIO)], [TestCase(**CASE)])

    def test_workbook_content_formatting_and_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            requirements, scenarios, cases = self.records()
            cases[0].title = "=1+1"
            target = Path(directory) / "nested" / "report.xlsx"
            export_to_excel(requirements, scenarios, cases, target)
            book = load_workbook(target)
            try:
                self.assertEqual(book.sheetnames, ["Requirements", "Scenarios", "Test Cases", "Traceability"])
                self.assertEqual(book["Requirements"]["A2"].value, "REQ-2")
                self.assertEqual(book["Test Cases"]["H2"].value, "1. Enter an eight-character password")
                self.assertEqual(book["Test Cases"]["D2"].data_type, "s")
                self.assertEqual(list(book["Traceability"].values)[1], ("REQ-2", "REQ001", "SC001", "TC001", "Functional", "=1+1", "Covered"))
                for sheet in book:
                    self.assertEqual(sheet.freeze_panes, "A2")
                    self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
                    self.assertTrue(sheet["A1"].font.bold)
                    self.assertTrue(sheet["A2"].alignment.wrap_text)
            finally:
                book.close()

    def test_uncovered_and_invalid_mappings(self):
        with tempfile.TemporaryDirectory() as directory:
            requirements, scenarios, cases = self.records()
            path = Path(directory) / "report.xlsx"
            export_to_excel(requirements, scenarios, [], path)
            book = load_workbook(path)
            self.assertEqual(book["Traceability"]["G2"].value, "Not covered")
            book.close()
            cases[0].scenario_id = "unknown"
            with self.assertRaises(ValueError):
                export_to_excel(requirements, scenarios, cases, path)

    def test_first_two_requirements_and_global_ids(self):
        with tempfile.TemporaryDirectory() as directory, contextlib.ExitStack() as stack:
            stack.enter_context(contextlib.redirect_stderr(io.StringIO()))
            extracted = {"requirements": [{"source_id": f"REQ-{i}", "requirement_text": f"Requirement {i}"} for i in range(1, 4)]}
            stack.enter_context(patch("services.requirement_extractor.call_llm", return_value=json.dumps(extracted)))
            analyze = stack.enter_context(patch("services.requirement_analyzer.call_llm", return_value=json.dumps(ANALYSIS)))
            stack.enter_context(patch("services.scenario_generator.call_llm", return_value=json.dumps({"scenarios": [SCENARIO]})))
            stack.enter_context(patch("services.test_case_generator.call_llm", return_value=json.dumps({"test_cases": [CASE]})))
            result = run_document_pipeline("sample_documents/login_brd.txt", Path(directory) / "report.xlsx")
            self.assertEqual(analyze.call_count, 2)
            self.assertEqual(result["processed_count"], 2)
            book = load_workbook(result["output_file"])
            rows = list(book["Traceability"].values)[1:]
            self.assertEqual([r[:4] for r in rows], [("REQ-1", "REQ001", "SC001", "TC001"), ("REQ-2", "REQ002", "SC002", "TC002")])
            book.close()
