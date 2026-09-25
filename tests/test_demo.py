import base64
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from api.demo import app, response_for
from services import demo_limits as limits
from services.llm_service import call_llm
from test_pipeline import ANALYSIS, SCENARIO, CASE


class DemoTests(unittest.TestCase):
    def mocks(self, scenarios=None, cases=None):
        stack = contextlib.ExitStack()
        payloads = [("requirement_analyzer", ANALYSIS),
                    ("scenario_generator", {"scenarios": scenarios or [SCENARIO]}),
                    ("test_case_generator", {"test_cases": cases or [CASE]})]
        for module, payload in payloads:
            stack.enter_context(patch(f"services.{module}.call_llm", return_value=json.dumps(payload)))
        return stack

    def test_demo_import_does_not_load_database(self):
        code = "import api.demo, sys; assert not any(x == 'runs' or x.startswith('runs.') for x in sys.modules)"
        result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_routes_are_stateless(self):
        with TestClient(app) as client:
            self.assertEqual(client.get("/health").json()["mode"], "demo")
            self.assertEqual(client.get("/docs").status_code, 200)
            self.assertEqual(client.get("/runs").status_code, 404)
            self.assertEqual(client.post("/runs", json={}).status_code, 404)

    def test_manual_roundtrip_preserves_ids_and_excel(self):
        with TestClient(app) as client, self.mocks():
            response = client.post("/generate", json={"requirement": "Password length 8 to 20"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["cache-control"], "no-store")
        report = response.json()
        result = report["results"][0]
        self.assertEqual(result["test_cases"][0]["requirement_id"], result["analysis"]["requirement_id"])
        workbook = load_workbook(io.BytesIO(base64.b64decode(report["excel_base64"])))
        self.assertEqual(workbook.sheetnames, ["Requirements", "Scenarios", "Test Cases", "Traceability"])
        self.assertEqual(workbook["Traceability"]["A2"].value, "SOURCE-001")
        workbook.close()

    def test_document_reports_processed_subset(self):
        extracted = {"requirements": [{"source_id": f"SOURCE-{i}", "requirement_text": "Password length 8 to 20"} for i in range(3)]}
        with TestClient(app) as client, self.mocks(), patch(
                "services.requirement_extractor.call_llm", return_value=json.dumps(extracted)):
            response = client.post("/generate-document", files={"file": ("brd.txt", b"Password length 8 to 20")})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["processed_count"], 2)
        self.assertEqual(response.json()["extracted_count"], 3)
        results = response.json()["results"]
        self.assertNotEqual(results[0]["analysis"]["requirement_id"], results[1]["analysis"]["requirement_id"])

    def test_inputs_rejected_before_ai(self):
        with TestClient(app) as client, patch("services.llm_service.OpenAI") as provider:
            for value in ["    ", "x" * (limits.MAX_TEXT_CHARS + 1)]:
                self.assertEqual(client.post("/generate", json={"requirement": value}).status_code, 422)
            for filename, content, expected in [("bad.csv", b"text", 400), ("empty.txt", b"", 400),
                    ("large.txt", b"x" * (limits.MAX_UPLOAD_BYTES + 1), 413),
                    ("long.txt", b"x" * (limits.MAX_DOCUMENT_CHARS + 1), 422)]:
                response = client.post("/generate-document", files={"file": (filename, content)})
                self.assertEqual(response.status_code, expected, response.text)
            provider.assert_not_called()

    def test_chunked_body_limit(self):
        with TestClient(app) as client:
            response = client.post("/generate-document", content=iter([b"x" * 2_100_000] * 2))
        self.assertEqual(response.status_code, 413)

    def test_excess_generated_scenarios_and_cases_are_not_silently_truncated(self):
        for scenarios, cases in [([SCENARIO] * 4, None), (None, [CASE] * 4)]:
            with TestClient(app) as client, self.mocks(scenarios, cases):
                response = client.post("/generate", json={"requirement": "Password length 8 to 20"})
            self.assertEqual(response.status_code, 422, response.text)

    def test_deadline_returns_504_and_resets_context(self):
        with TestClient(app) as client, patch.object(limits, "DEADLINE_SECONDS", -1), self.mocks():
            response = client.post("/generate", json={"requirement": "Password length 8 to 20"})
        self.assertEqual(response.status_code, 504)
        self.assertFalse(limits.active())

    def test_provider_budget_and_token_limit(self):
        with patch.dict(os.environ, {"OPENROUTER_API_KEY": "test-key"}), patch("services.llm_service.OpenAI") as provider:
            provider.return_value.__enter__.return_value.chat.completions.create.return_value.choices = []
            with limits.demo_budget():
                for _ in range(12):
                    call_llm("test")
                with self.assertRaises(limits.DemoLimitError):
                    call_llm("test")
            self.assertEqual(provider.call_count, 12)
            self.assertLessEqual(provider.call_args.kwargs["timeout"], 60)
            self.assertEqual(provider.return_value.__enter__.return_value.chat.completions.create.call_args.kwargs["max_tokens"], 4000)

    def test_oversize_response_rejected(self):
        with limits.demo_budget(), self.assertRaises(limits.DemoLimitError):
            response_for({"excel_data": b"x" * limits.MAX_RESPONSE_BYTES})

    def test_demo_ui_session_and_no_history(self):
        with TestClient(app) as client, self.mocks():
            report = client.post("/generate", json={"requirement": "Password length 8 to 20"}).json()
        report["excel_data"] = base64.b64decode(report.pop("excel_base64"))
        with patch.dict(os.environ, {"TESTGEN_MODE": "demo"}), \
                patch("ui.api_client.get_history") as history, \
                patch("ui.api_client.generate_demo", return_value=report) as generate:
            at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "ui/streamlit_app.py"), default_timeout=15).run()
            at.button[0].click().run()
            at.button[1].click().run()
            self.assertFalse(at.exception)
            self.assertEqual(len(at.metric), 5)
            at.run()
            generate.assert_called_once()
            history.assert_not_called()
            self.assertEqual(len(at.metric), 5)
            at.text_area[0].set_value("Different requirement text").run()
            self.assertEqual(len(at.metric), 0)
