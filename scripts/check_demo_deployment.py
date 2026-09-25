"""Check the API with real Excel IO; --live makes one real sample generation.

Run from the repository root. Output contains timing and counts, never credentials.
"""
import argparse
import base64
from contextlib import ExitStack
import io
import json
import os
from pathlib import Path
import sys
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from api.demo import app

SAMPLE = "Password length must be between 8 and 20 characters, inclusive."
ANALYSIS = dict(requirement_id="REQ001", actor="User", functionality="Create password",
                business_rules=[SAMPLE], constraints=["8 to 20 characters"],
                boundary_constraints=[], missing_information=[], assumptions=[],
                clarification_questions=[])
SCENARIO = dict(scenario_id="SC001", requirement_id="REQ001", title="Valid length",
                scenario_type="Functional", description="Accept valid password length")
CASE = dict(test_case_id="TC001", requirement_id="REQ001", scenario_id="SC001",
            title="Accept eight characters", test_type="Functional", priority="High",
            preconditions=[], steps=["Enter an eight-character password"],
            test_data={"password": "abcdefgh"}, expected_result="Length is accepted")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true", help="Call OpenRouter with the synthetic sample")
    args = parser.parse_args()
    if args.live and not os.getenv("OPENROUTER_API_KEY", "").strip():
        print("Live check requires OPENROUTER_API_KEY in local .env or the process environment.")
        return 2
    started = time.monotonic()
    try:
        with ExitStack() as stack:
            if not args.live:
                for module, payload in [("requirement_analyzer", ANALYSIS),
                                        ("scenario_generator", {"scenarios": [SCENARIO]}),
                                        ("test_case_generator", {"test_cases": [CASE]})]:
                    stack.enter_context(patch(f"services.{module}.call_llm", return_value=json.dumps(payload)))
            client = stack.enter_context(TestClient(app))
            assert client.get("/health").json() == {"status": "healthy", "mode": "demo"}
            assert client.get("/docs").status_code == 200
            assert client.get("/runs").status_code == 404
            response = client.post("/generate", json={"requirement": SAMPLE})
            if response.status_code != 200:
                print(json.dumps({"mode": "live" if args.live else "mocked", "status": response.status_code,
                                  "detail": response.json().get("detail", "Request failed"),
                                  "elapsed_seconds": round(time.monotonic() - started, 2)}))
                return 1
            report = response.json()
            workbook = load_workbook(io.BytesIO(base64.b64decode(report["excel_base64"], validate=True)))
            assert workbook.sheetnames == ["Requirements", "Scenarios", "Test Cases", "Traceability"]
            assert workbook["Traceability"].max_row > 1
            workbook.close()
            assert not any(name == "runs" or name.startswith("runs.") for name in sys.modules)
            count = sum(len(result["test_cases"]) for result in report["results"])
            assert count > 0, "No accepted test cases"
            print(json.dumps({"mode": "live" if args.live else "mocked", "status": "passed",
                              "elapsed_seconds": round(time.monotonic() - started, 2),
                              "accepted_cases": count, "response_bytes": len(response.content)}))
            return 0
    except Exception as exc:
        print(json.dumps({"status": "failed", "error_type": type(exc).__name__,
                          "elapsed_seconds": round(time.monotonic() - started, 2)}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
