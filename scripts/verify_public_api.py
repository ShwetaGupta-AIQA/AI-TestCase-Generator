"""Smoke-check public design/export endpoints with synthetic input; no target API execution."""
import json
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree

import httpx
import yaml
from openpyxl import load_workbook


def main():
    spec = {"openapi": "3.0.3", "info": {"title": "Synthetic deployment check", "version": "1"},
            "paths": {"/items": {"post": {"requestBody": {"content": {"application/json": {
                "schema": {"type": "object", "required": ["quantity"], "properties": {
                    "quantity": {"type": "integer", "minimum": 1, "maximum": 5}}}}}},
                "responses": {"201": {"description": "Created"}}}}}}
    with httpx.Client(base_url="https://testgen-ai-api.onrender.com", timeout=90) as client:
        def post(path, payload):
            response = client.post(path, json=payload)
            response.raise_for_status()
            return response
        for label, text in [("JSON", json.dumps(spec)), ("YAML", yaml.safe_dump(spec))]:
            parsed = post("/api-testing/spec/parse", {"spec_text": text}).json()
            endpoint = parsed["endpoints"][0]["contract"]
            endpoint["responses"] = [{"status_code": "202", "description": "Reviewed"}]
            plan = post("/api-testing/design/spec", {"spec_text": text, "reviewed_endpoint": endpoint}).json()["plan"]
            assert plan["test_cases"][0]["expected_status"] == "202"
            assert isinstance(plan["test_cases"][0]["request_body"]["quantity"], int)
            assert any(case["test_type"] == "Boundary/Validation" for case in plan["test_cases"])
            print(label, "import, review, typed payload, boundary generation: PASS", flush=True)
        plan["test_cases"] = plan["test_cases"][:1]
        plan["coverage_summary"] = {"Functional": 1}
        for kind in ["excel", "postman", "pytest", "soapui"]:
            data = post("/api-testing/export/" + kind, {"plan": plan}).content
            if kind == "excel":
                book = load_workbook(BytesIO(data)); assert book["API Test Cases"].max_row == 2; book.close()
            elif kind == "postman":
                assert len(json.loads(data)["item"]) == 1
            elif kind == "pytest":
                with ZipFile(BytesIO(data)) as archive:
                    compile(archive.read("test_api_contract.py"), "exported.py", "exec")
            else:
                ElementTree.fromstring(data)
            print(kind, "selected-plan export: PASS", flush=True)
        html = "<h1>Example</h1><p>GET /items/{id}</p><p>Response status 200</p>"
        parsed = post("/api-testing/html/parse", {"html_text": html}).json()
        assert parsed["base_url"] == ""
        post("/api-testing/design/html", {"html_text": html, "endpoint_key": parsed["endpoints"][0]["key"]})
        print("HTML import and generation: PASS", flush=True)
        bad = client.post("/api-testing/spec/parse", json={"spec_text": "openapi: [broken"})
        assert 400 <= bad.status_code < 500
        print("Invalid input handling: PASS", flush=True)
    response = httpx.get("https://testgen-ai-demo.streamlit.app/", follow_redirects=True, timeout=60)
    response.raise_for_status()
    print("Streamlit HTTP:", response.status_code, "(interactive UI not verified)", flush=True)


if __name__ == "__main__":
    main()
