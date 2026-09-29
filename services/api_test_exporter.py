import io
import json
import re
import zipfile
from urllib.parse import urlencode
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.api_testing import ApiTestCase, ApiTestPlan


def _safe_name(value: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9_]+", "_", value).strip("_").lower()
    return clean or "api_test"


def _url(plan: ApiTestPlan, case: ApiTestCase) -> str:
    return (plan.base_url.rstrip("/") + "/" + case.path.lstrip("/")) if plan.base_url else case.path


def export_api_plan_excel(plan: ApiTestPlan) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    specs = [
        ("Summary", ["Field", "Value"], [28, 80]),
        ("API Test Cases", ["Test Case ID", "Title", "Type", "Priority", "Method", "Path", "Expected Status", "Objective", "Steps", "Expected Result"], [16, 45, 20, 12, 10, 45, 18, 60, 80, 70]),
        ("Request Data", ["Test Case ID", "Headers", "Query Params", "Path Params", "Body"], [16, 60, 60, 60, 70]),
        ("Traceability", ["Endpoint", "Coverage Type", "Test Case ID", "Coverage Status"], [40, 24, 16, 18]),
    ]
    for name, headers, widths in specs:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    summary = workbook["Summary"]
    endpoint = f"{plan.endpoint.method} {plan.endpoint.path}"
    for row in [
        ("Contract", plan.contract_title),
        ("Base URL", plan.base_url),
        ("Endpoint", endpoint),
        ("Generated Cases", len(plan.test_cases)),
        ("Coverage Summary", ", ".join(f"{name}: {count}" for name, count in plan.coverage_summary.items())),
        ("Retrieved Context", "\n".join(plan.retrieved_context)),
    ]:
        summary.append(row)

    for case in plan.test_cases:
        workbook["API Test Cases"].append([
            case.test_case_id, case.title, case.test_type, case.priority, case.method, case.path,
            case.expected_status, case.objective, "\n".join(f"{i}. {step}" for i, step in enumerate(case.steps, 1)),
            case.expected_result,
        ])
        workbook["Request Data"].append([
            case.test_case_id,
            json.dumps(case.request_headers, indent=2),
            json.dumps(case.query_params, indent=2),
            json.dumps(case.path_params, indent=2),
            json.dumps(case.request_body, indent=2),
        ])
        workbook["Traceability"].append([endpoint, case.test_type, case.test_case_id, "Designed; not executed"])

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 36
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203864")

    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def export_postman_collection(plan: ApiTestPlan) -> bytes:
    collection = {
        "info": {
            "name": f"TestGen API Tests - {plan.endpoint.method} {plan.endpoint.path}",
            "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
        },
        "item": [],
        "variable": [{"key": "base_url", "value": plan.base_url.rstrip("/")}],
    }
    for case in plan.test_cases:
        collection["item"].append({
            "name": f"{case.test_case_id} - {case.title}",
            "request": {
                "method": case.method,
                "header": [{"key": key, "value": value} for key, value in case.request_headers.items()],
                "url": "{{base_url}}/" + case.path.lstrip("/") + ("?" + urlencode(case.query_params) if case.query_params else ""),
                "body": {"mode": "raw", "raw": json.dumps(case.request_body, indent=2),
                         "options": {"raw": {"language": "json"}}} if case.request_body else None,
            },
            "event": [{"listen": "test", "script": {"type": "text/javascript", "exec": [
                f"pm.test('Expected status {case.expected_status}', function () {{",
                f"    pm.response.to.have.status({case.expected_status});" if case.expected_status.isdigit() else "    throw new Error('Expected status needs confirmation');",
                "});",
            ]}}],
        })
        item = collection["item"][-1]
        item["request"]["description"] = case.expected_result + "\nAutomated assertions check HTTP status only."
        if item["request"]["body"] is None:
            del item["request"]["body"]
        elif not any(key.lower() == "content-type" for key in case.request_headers):
            item["request"]["header"].append({"key": "Content-Type", "value": "application/json"})
        if not case.expected_status.isdigit():
            item["event"] = [{"listen": "prerequest", "script": {"type": "text/javascript",
                "exec": ["// Confirm the expected status in TestGen and export again.", "pm.execution.skipRequest();"]}}]
    return json.dumps(collection, indent=2).encode("utf-8")


def export_pytest_suite(plan: ApiTestPlan) -> bytes:
    lines = [
        "import os",
        "",
        "import requests",
        "import pytest",
        "",
        f"BASE_URL = os.getenv('API_BASE_URL', {plan.base_url!r}).rstrip('/')",
        "",
        "",
        "def _url(path):",
        "    return BASE_URL + '/' + path.lstrip('/') if BASE_URL else path",
        "",
    ]
    for case in plan.test_cases:
        lines.extend([
            f"def test_{_safe_name(case.test_case_id + '_' + case.title)}():",
            "    pass" if case.expected_status.isdigit() else "    pytest.skip('Expected status needs confirmation')",
            f"    response = requests.request({case.method!r}, _url({case.path!r}),",
            f"        headers={case.request_headers!r},",
            f"        params={case.query_params!r},",
            f"        json={case.request_body!r} if {bool(case.request_body)!r} else None, timeout=30)",
            f"    assert response.status_code == {int(case.expected_status) if case.expected_status.isdigit() else 200}",
            "",
        ])
    content = "\n".join(lines).encode("utf-8")
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("test_api_contract.py", content)
        archive.writestr("requirements.txt", "pytest>=8,<10\nrequests>=2.32,<3\n")
        archive.writestr("README.md", """# Run the generated API tests

1. Install dependencies: `python -m pip install -r requirements.txt`
2. Configure an authorized Mock, Local, Test or Staging API base URL.
   PowerShell: `$env:API_BASE_URL = 'http://localhost:8000'`
   Bash: `export API_BASE_URL=http://localhost:8000`
3. Review requests and replace credential/test-data placeholders in the test file.
4. Run: `python -m pytest test_api_contract.py -v`

Cases with unconfirmed expected statuses are skipped before sending a request.
The generated assertions currently check HTTP status only. Expected-result prose,
schema, database, idempotency and business-rule checks require additional assertions
and setup. A passing status assertion does not establish those outcomes.
""")
    return stream.getvalue()


def export_soapui_project(plan: ApiTestPlan) -> bytes:
    con = "http://eviware.com/soapui/config"
    ET.register_namespace("con", con)
    q = lambda name: f"{{{con}}}{name}"
    project = ET.Element("con:soapui-project", {
        "name": f"TestGen API Tests - {plan.contract_title}",
        "id": "testgen-api-project",
        "activeEnvironment": "Default",
        "soapui-version": "5.7.2",
    })
    ET.SubElement(project, q("settings"))
    ET.SubElement(project, q("properties"))
    ET.SubElement(project, q("wssContainer"))
    interface = ET.SubElement(project, q("interface"), {
        "xsi:type": "con:RestService",
        "xmlns:xsi": "http://www.w3.org/2001/XMLSchema-instance",
        "id": "testgen-rest-interface",
        "name": f"{plan.endpoint.method} {plan.endpoint.path}",
        "type": "rest",
    })
    ET.SubElement(interface, q("settings"))
    definition = ET.SubElement(interface, q("definitionCache"), {"type": "TEXT"})
    ET.SubElement(definition, q("part")).text = json.dumps({
        "contract": plan.contract_title,
        "endpoint": f"{plan.endpoint.method} {plan.endpoint.path}",
        "generatedBy": "TestGen AI",
    }, indent=2)
    ET.SubElement(ET.SubElement(interface, q("endpoints")), q("endpoint")).text = plan.base_url or "http://localhost"
    resource = ET.SubElement(interface, q("resource"), {
        "name": plan.endpoint.path,
        "path": plan.endpoint.path,
        "id": "testgen-resource",
    })
    ET.SubElement(resource, q("parameters"))
    for case in plan.test_cases:
        method = ET.SubElement(resource, q("method"), {"name": case.test_case_id, "method": case.method})
        parameters = ET.SubElement(method, q("parameters"))
        for style, values in [("HEADER", case.request_headers), ("QUERY", case.query_params), ("TEMPLATE", case.path_params)]:
            for name, value in values.items():
                parameter = ET.SubElement(parameters, q("parameter"))
                ET.SubElement(parameter, q("name")).text = name
                ET.SubElement(parameter, q("value")).text = str(value)
                ET.SubElement(parameter, q("style")).text = style
                ET.SubElement(parameter, q("default")).text = str(value)
        request = ET.SubElement(method, q("request"), {
            "name": f"{case.test_case_id} - {case.title}",
            "mediaType": "application/json",
            "postQueryString": "false",
            "id": f"request-{case.test_case_id}",
        })
        ET.SubElement(request, q("settings"))
        ET.SubElement(request, q("encoding")).text = "UTF-8"
        ET.SubElement(request, q("endpoint")).text = plan.base_url or "http://localhost"
        ET.SubElement(request, q("request")).text = json.dumps(case.request_body, indent=2) if case.request_body else ""
        assertions = ET.SubElement(request, q("assertion"), {
            "type": "Valid HTTP Status Codes",
            "name": f"Expected status {case.expected_status}",
            "id": f"assertion-{case.test_case_id}",
        })
        ET.SubElement(assertions, q("configuration")).text = case.expected_status if case.expected_status.isdigit() else ""

    test_suite = ET.SubElement(project, q("testSuite"), {"name": f"{plan.endpoint.method} {plan.endpoint.path}"})
    ET.SubElement(test_suite, q("settings"))
    for case in plan.test_cases:
        soap_case = ET.SubElement(test_suite, q("testCase"), {"name": f"{case.test_case_id} - {case.title}"})
        ET.SubElement(soap_case, q("settings"))
        step = ET.SubElement(soap_case, q("testStep"), {"type": "restrequest", "name": case.test_case_id})
        config = ET.SubElement(step, q("config"), {
            "service": f"{plan.endpoint.method} {plan.endpoint.path}",
            "resourcePath": plan.endpoint.path,
            "methodName": case.test_case_id,
            "xsi:type": "con:RestRequestStep",
            "xmlns:xsi": "http://www.w3.org/2001/XMLSchema-instance",
        })
        rest_request = ET.SubElement(config, q("restRequest"), {
            "name": case.test_case_id,
            "mediaType": "application/json",
            "postQueryString": "false",
        })
        ET.SubElement(rest_request, q("settings"))
        ET.SubElement(rest_request, q("encoding")).text = "UTF-8"
        ET.SubElement(rest_request, q("endpoint")).text = plan.base_url or "http://localhost"
        ET.SubElement(rest_request, q("request")).text = json.dumps(case.request_body, indent=2)
        if case.expected_status.isdigit():
            assertion = ET.SubElement(rest_request, q("assertion"), {"type": "Valid HTTP Status Codes", "name": "Expected status"})
            ET.SubElement(ET.SubElement(assertion, q("configuration")), "codes").text = case.expected_status
        else:
            soap_case.set("disabled", "true")
        ET.SubElement(rest_request, q("credentials"))
        ET.SubElement(rest_request, q("parameters"))
        ET.SubElement(soap_case, q("properties"))
    ET.SubElement(project, q("mockService"), {"name": "Generated Mock Placeholder"})
    ET.SubElement(project, q("environments"))
    ET.SubElement(project, q("authRepository"))
    ET.SubElement(project, q("tags"))
    ET.indent(project)
    return ET.tostring(project, encoding="utf-8", xml_declaration=True)
