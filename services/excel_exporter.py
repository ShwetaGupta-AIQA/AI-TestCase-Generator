"""Export validated QA records with source-to-case traceability."""
from pathlib import Path
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_to_excel(requirements, scenarios, test_cases, output_file="TestGen_Output.xlsx"):
    source_mapping = {item["analysis"].requirement_id: item["source"].source_id for item in requirements}
    scenario_mapping = {s.scenario_id: s.requirement_id for s in scenarios}
    if len(source_mapping) != len(requirements) or len(scenario_mapping) != len(scenarios):
        raise ValueError("Export requires unique requirement and scenario IDs.")
    if len({c.test_case_id for c in test_cases}) != len(test_cases):
        raise ValueError("Export requires unique test case IDs.")
    for scenario in scenarios:
        if scenario.requirement_id not in source_mapping:
            raise ValueError("Scenario has no source requirement mapping.")
    for case in test_cases:
        if case.requirement_id not in source_mapping or scenario_mapping.get(case.scenario_id) != case.requirement_id:
            raise ValueError("Test case has an invalid traceability mapping.")

    workbook = Workbook()
    workbook.remove(workbook.active)
    specs = [
        ("Requirements", ["Source ID", "Requirement ID", "Requirement Text", "Actor", "Functionality", "Business Rules", "Constraints", "Missing Information"], [15, 18, 50, 20, 35, 50, 50, 50]),
        ("Scenarios", ["Scenario ID", "Requirement ID", "Type", "Title", "Description"], [15, 18, 15, 45, 60]),
        ("Test Cases", ["Test Case ID", "Requirement ID", "Scenario ID", "Title", "Test Type", "Priority", "Preconditions", "Steps", "Test Data", "Expected Result"], [15, 18, 15, 45, 15, 12, 40, 70, 40, 60]),
        ("Traceability", ["Source Requirement", "Requirement ID", "Scenario ID", "Test Case ID", "Test Type", "Test Case Title", "Coverage Status"], [24, 18, 15, 18, 15, 45, 20]),
    ]
    for name, headers, widths in specs:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    for item in requirements:
        source, analysis = item["source"], item["analysis"]
        workbook["Requirements"].append([
            source.source_id, analysis.requirement_id, source.requirement_text,
            analysis.actor, analysis.functionality, "\n".join(analysis.business_rules),
            "\n".join(analysis.constraints), "\n".join(analysis.missing_information)])
    for scenario in scenarios:
        workbook["Scenarios"].append([scenario.scenario_id, scenario.requirement_id,
            scenario.scenario_type, scenario.title, scenario.description])
    for case in test_cases:
        workbook["Test Cases"].append([case.test_case_id, case.requirement_id,
            case.scenario_id, case.title, case.test_type, case.priority,
            "\n".join(case.preconditions),
            "\n".join(f"{i}. {step}" for i, step in enumerate(case.steps, 1)),
            "\n".join(f"{key}: {value}" for key, value in case.test_data.items()),
            case.expected_result])
        workbook["Traceability"].append([source_mapping[case.requirement_id],
            case.requirement_id, case.scenario_id, case.test_case_id,
            case.test_type, case.title, "Covered"])
    # Show gaps rather than silently hiding scenarios without accepted cases.
    covered = {case.scenario_id for case in test_cases}
    for scenario in scenarios:
        if scenario.scenario_id not in covered:
            workbook["Traceability"].append([source_mapping[scenario.requirement_id],
                scenario.requirement_id, scenario.scenario_id, "", "", "", "Not covered"])
    for requirement_id, source_id in source_mapping.items():
        if requirement_id not in scenario_mapping.values():
            workbook["Traceability"].append([source_id, requirement_id, "", "", "", "", "Not covered"])

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            line_count = 1
            for cell in row:
                # AI/source text must remain literal even when it starts with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                width = sheet.column_dimensions[cell.column_letter].width
                line_count = max(line_count, sum(max(1, math.ceil(len(line) / max(1, width - 2)))
                    for line in str(cell.value or "").split("\n")))
            sheet.row_dimensions[row[0].row].height = min(409, max(30, 16 * line_count + 8))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203864")
    path = Path(output_file)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return str(path)
